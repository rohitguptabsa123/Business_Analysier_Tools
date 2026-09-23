"""
Generate a PowerPoint presentation for DataFlow Analyzer.

Output: demo_output/DataFlow_Analyzer_Presentation.pptx

Run:  py make_presentation.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN as PPAlign, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "demo_output")
os.makedirs(OUT, exist_ok=True)

# ── Colors ──
C_INPUT   = RGBColor(0x4C, 0xAF, 0x50)
C_PROC    = RGBColor(0x21, 0x96, 0xF3)
C_ANAL    = RGBColor(0x9C, 0x27, 0xB0)
C_CHART   = RGBColor(0xFF, 0x98, 0x00)
C_EXP     = RGBColor(0xF4, 0x43, 0x36)
C_UI      = RGBColor(0x60, 0x7D, 0x8B)
C_BG      = RGBColor(0x1A, 0x1A, 0x2E)
C_PANEL   = RGBColor(0x16, 0x21, 0x3E)
C_ACCENT  = RGBColor(0x0F, 0x34, 0x60)
C_TEXT    = RGBColor(0xE0, 0xE0, 0xE0)
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_HIGHLIGHT = RGBColor(0xE9, 0x45, 0x60)
C_DARK    = RGBColor(0x0A, 0x16, 0x28)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_bg(slide, color=C_BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, color, alpha=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Inches(left), Inches(top),
                                    Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_text(slide, left, top, width, height, text, size=18, color=C_WHITE,
             bold=False, align=PPAlign.LEFT, font="Segoe UI"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font
    p.alignment = align
    return txBox

def add_bullet_list(slide, left, top, width, height, items, size=16, color=C_TEXT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)
    return txBox

def add_feature_card(slide, left, top, width, height, title, desc, color):
    add_shape(slide, left, top, width, height, C_PANEL)
    # Color bar at top
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(left), Inches(top),
                                  Inches(width), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    add_text(slide, left + 0.2, top + 0.2, width - 0.4, 0.5, title,
             size=16, color=color, bold=True)
    add_text(slide, left + 0.2, top + 0.7, width - 0.4, height - 0.9, desc,
             size=12, color=C_TEXT)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 1: Title
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide)
add_text(slide, 1, 1.5, 11.3, 1.2, "DataFlow Analyzer",
         size=48, color=C_WHITE, bold=True, align=PPAlign.CENTER)
add_text(slide, 1, 2.8, 11.3, 0.8, "Desktop Data Analysis Application",
         size=24, color=C_CHART, align=PPAlign.CENTER)
add_text(slide, 1, 3.8, 11.3, 0.5, "Import | Filter | Pivot | Chart | Analyze | Export",
         size=16, color=C_TEXT, align=PPAlign.CENTER)
# Feature badges
features = [("CSV/Excel", C_INPUT), ("Filter & Pivot", C_PROC),
            ("Charts", C_CHART), ("Regression", C_ANAL), ("K-means", C_INPUT), ("Export", C_EXP)]
for i, (feat, color) in enumerate(features):
    fx = 1.5 + i * 1.8
    add_shape(slide, fx, 4.8, 1.6, 0.5, color)
    add_text(slide, fx, 4.85, 1.6, 0.4, feat, size=11, color=C_WHITE, bold=True,
             align=PPAlign.CENTER)
add_text(slide, 1, 6.2, 11.3, 0.5, "Built with Python, Tkinter, pandas & matplotlib",
         size=12, color=C_TEXT, align=PPAlign.CENTER)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 2: Overview
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.4, 11, 0.8, "Application Overview", size=36, color=C_WHITE, bold=True)
add_text(slide, 0.8, 1.1, 11, 0.5, "A standalone desktop tool for everyday data analysis",
         size=16, color=C_CHART)

add_bullet_list(slide, 0.8, 2.0, 5.5, 4.5, [
    "\u25cf  Import CSV, TSV, and Excel files (up to 500,000 rows)",
    "\u25cf  Non-blocking background thread for file loading",
    "\u25cf  Paginated data table (1,000 rows per page)",
    "\u25cf  Interactive filtering with 7 operators",
    "\u25cf  Pivot tables with 7 aggregation functions",
    "\u25cf  Bar, line, and scatter charts (matplotlib)",
    "\u25cf  Linear regression (NumPy OLS)",
    "\u25cf  K-means clustering (Lloyd's algorithm)",
    "\u25cf  Export to formatted Excel (xlsxwriter)",
    "\u25cf  Export to PDF reports (reportlab)",
], size=15, color=C_TEXT)

# Right side: architecture summary
add_shape(slide, 7, 2.0, 5.5, 4.8, C_PANEL)
add_text(slide, 7.2, 2.2, 5, 0.5, "Architecture", size=18, color=C_CHART, bold=True)
layers = [
    ("Data Input Layer", "CSV / Excel / TSV loader", C_INPUT),
    ("GUI Layer", "Tkinter window, 5 tabs, toolbar", C_UI),
    ("Processing Layer", "pandas DataFrame, filter, pivot", C_PROC),
    ("Analysis & ML Layer", "Statistics, regression, k-means", C_ANAL),
    ("Visualization Layer", "matplotlib charts (bar/line/scatter)", C_CHART),
    ("Export Layer", "Excel (xlsxwriter) + PDF (reportlab)", C_EXP),
]
for i, (title, desc, color) in enumerate(layers):
    y = 2.8 + i * 0.65
    add_shape(slide, 7.2, y, 0.15, 0.5, color)
    add_text(slide, 7.5, y, 4.8, 0.3, title, size=12, color=color, bold=True)
    add_text(slide, 7.5, y + 0.28, 4.8, 0.3, desc, size=10, color=C_TEXT)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 3: Key Features (6 cards)
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.3, 11, 0.8, "Key Features", size=36, color=C_WHITE, bold=True)

cards = [
    (0.5, 1.4, 3.8, 2.6, "Data Import",
     "Load CSV, TSV, and Excel files in a background thread. Handles up to 500K rows without blocking the UI.",
     C_INPUT),
    (4.7, 1.4, 3.8, 2.6, "Filter & Pivot",
     "7 filter operators (==, !=, >, >=, <, <=, contains). Pivot tables with mean, sum, count, min, max, median, std.",
     C_PROC),
    (8.9, 1.4, 3.8, 2.6, "Charts",
     "Bar, line, and scatter charts rendered with matplotlib. Animated fade-in on draw. Export-ready figures.",
     C_CHART),
    (0.5, 4.3, 3.8, 2.6, "Regression",
     "Linear regression using NumPy OLS. Reports slope, intercept, and R\u00b2. Works on any numeric columns.",
     C_ANAL),
    (4.7, 4.3, 3.8, 2.6, "K-Means Clustering",
     "Lloyd's algorithm implementation. Configurable k (2-10). Assigns cluster labels to every row.",
     C_INPUT),
    (8.9, 4.3, 3.8, 2.6, "Export",
     "Excel: formatted headers, autofilter, frozen panes. PDF: summary stats, data preview, embedded charts.",
     C_EXP),
]
for (lx, ly, lw, lh, title, desc, color) in cards:
    add_feature_card(slide, lx, ly, lw, lh, title, desc, color)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 4: UI Animations
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.3, 11, 0.8, "UI Animations", size=36, color=C_WHITE, bold=True)
add_text(slide, 0.8, 1.0, 11, 0.5, "Rich animations throughout the application for a polished experience",
         size=16, color=C_CHART)

animations = [
    ("Typewriter Status Bar", "Status messages appear character-by-character with variable speed",
     "12-25ms per character, auto-cancels on new message", C_PROC),
    ("Pulsing Status Indicator", "Color-coded dot pulses continuously (green/blue/red)",
     "Sine-wave pulse, 3-ring glow halo, 60ms refresh", C_INPUT),
    ("Loading Spinner", "8-segment rotating arc spinner during file loading",
     "15\u00b0 per frame, 50ms interval, blue fade gradient", C_CHART),
    ("Staggered Row Reveal", "Data table rows appear in batches of 50 with 15ms delay",
     "Smooth fill-in effect, cancellable on page change", C_ANAL),
    ("Chart Fade-In", "Charts progressively fade from 0% to 100% opacity",
     "All artists animated, 30ms per frame, ~7 frames", C_EXP),
    ("Progress Bar", "Indeterminate progress bar during all operations",
     "Filter, pivot, stats, regression, k-means, export", C_UI),
]
for i, (title, desc, detail, color) in enumerate(animations):
    y = 1.8 + i * 0.85
    add_shape(slide, 0.8, y, 0.15, 0.7, color)
    add_text(slide, 1.1, y, 4, 0.35, title, size=14, color=color, bold=True)
    add_text(slide, 1.1, y + 0.32, 5, 0.35, desc, size=11, color=C_TEXT)
    add_text(slide, 6.5, y + 0.15, 6, 0.4, detail, size=10, color=C_CHART)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 5: Technology Stack
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.3, 11, 0.8, "Technology Stack", size=36, color=C_WHITE, bold=True)

tech = [
    ("Python 3.12+", "Core language, type hints, modern syntax", C_PROC),
    ("Tkinter / ttk", "GUI framework - window, tabs, treeview, toolbar", C_UI),
    ("pandas", "DataFrame operations, filtering, pivot tables, I/O", C_PROC),
    ("NumPy", "Numerical computing, OLS regression, array ops", C_ANAL),
    ("matplotlib", "Chart rendering (bar, line, scatter), embedded in Tk", C_CHART),
    ("xlsxwriter", "Excel export with formatting, autofilter, frozen panes", C_EXP),
    ("reportlab", "PDF report generation with tables, charts, statistics", C_EXP),
    ("unittest", "Unit testing framework (12 tests, all passing)", C_INPUT),
]
for i, (name, desc, color) in enumerate(tech):
    y = 1.4 + i * 0.7
    add_shape(slide, 0.8, y, 3, 0.55, color)
    add_text(slide, 0.8, y + 0.08, 3, 0.4, name, size=13, color=C_WHITE, bold=True,
             align=PPAlign.CENTER)
    add_text(slide, 4.2, y + 0.1, 8, 0.4, desc, size=13, color=C_TEXT)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 6: Project Structure
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.3, 11, 0.8, "Project Structure", size=36, color=C_WHITE, bold=True)

structure = [
    ("dataflow/", "Main package", C_CHART, True),
    ("  \u251c\u2500 __init__.py", "Package init, exports main()", C_TEXT, False),
    ("  \u251c\u2500 app.py", "Tkinter GUI - DataFlowApp class, all UI + animations", C_PROC, False),
    ("  \u251c\u2500 data_loader.py", "Async file loader (CSV/Excel/TSV)", C_INPUT, False),
    ("  \u251c\u2500 analysis.py", "Filter, pivot, describe, correlation, regression, k-means", C_ANAL, False),
    ("  \u251c\u2500 charts.py", "Bar/line/scatter chart rendering", C_CHART, False),
    ("  \u2514\u2500 exporter.py", "Excel (xlsxwriter) + PDF (reportlab) export", C_EXP, False),
    ("tests/", "Unit tests (12 tests, all passing)", C_INPUT, True),
    ("sample_data/", "Sample CSV files for demo", C_INPUT, True),
    ("run.py", "Entry point: py run.py", C_WHITE, False),
    ("make_flow_diagram.py", "Static architecture flowchart generator", C_CHART, False),
    ("make_animated_flow.py", "Animated GIF flow diagram generator", C_CHART, False),
    ("make_demo_video.py", "2-minute demo video generator", C_CHART, False),
    ("make_presentation.py", "This PowerPoint generator", C_CHART, False),
]
for i, (name, desc, color, is_header) in enumerate(structures := structure):
    y = 1.2 + i * 0.42
    sz = 14 if is_header else 11
    add_text(slide, 0.8, y, 4, 0.35, name, size=sz, color=color,
             bold=is_header, font="Consolas")
    add_text(slide, 5.0, y, 7.5, 0.35, desc, size=sz, color=C_TEXT)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 7: Demo Walkthrough
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.3, 11, 0.8, "Demo Walkthrough", size=36, color=C_WHITE, bold=True)
add_text(slide, 0.8, 1.0, 11, 0.5, "2-minute video showcasing the full workflow",
         size=16, color=C_CHART)

steps = [
    ("1. Open File", "Click Open File or press Ctrl+O. Select a CSV/Excel file. Loading spinner + progress bar animate while data loads in a background thread.", C_INPUT),
    ("2. Data Table", "Rows appear with staggered batch reveal (50 rows per 15ms). Paginated at 1,000 rows/page. Status bar types out row count.", C_PROC),
    ("3. Filter Data", "Select column, operator, and value. Progress bar animates. Filtered results appear with staggered reveal.", C_PROC),
    ("4. Charts", "Choose bar, line, or scatter. Chart fades in from 0% to 100% opacity. All artists animate together.", C_CHART),
    ("5. Analysis & ML", "Summary statistics, correlation matrix, linear regression (R\u00b2), and K-means clustering. Progress bar for each.", C_ANAL),
    ("6. Export", "Export to formatted Excel (autofilter, frozen panes) or PDF report (tables + charts + stats). Progress bar during export.", C_EXP),
]
for i, (title, desc, color) in enumerate(steps):
    y = 1.7 + i * 0.9
    add_shape(slide, 0.8, y, 0.15, 0.75, color)
    add_text(slide, 1.1, y, 3.5, 0.35, title, size=14, color=color, bold=True)
    add_text(slide, 1.1, y + 0.32, 11, 0.45, desc, size=11, color=C_TEXT)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 8: Testing & Quality
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.3, 11, 0.8, "Testing & Quality", size=36, color=C_WHITE, bold=True)

# Left: test results
add_shape(slide, 0.8, 1.3, 5.5, 5.5, C_PANEL)
add_text(slide, 1.0, 1.5, 5, 0.5, "Unit Tests", size=18, color=C_INPUT, bold=True)
add_bullet_list(slide, 1.0, 2.2, 5, 4, [
    "\u2713  test_filter_contains - string contains filter",
    "\u2713  test_filter_equality - == operator",
    "\u2713  test_filter_numeric_gt - > operator",
    "\u2713  test_kmeans_clusters - 3 clusters assigned",
    "\u2713  test_pivot_mean - mean aggregation",
    "\u2713  test_regression_perfect_fit - R\u00b2 = 1.0",
    "\u2713  test_correlation - 4x4 matrix",
    "\u2713  test_describe - 7 stats columns",
    "\u2713  test_load_csv - 2000 rows loaded",
    "\u2713  test_unsupported_ext - raises ValueError",
    "\u2713  test_export_excel - .xlsx created",
    "\u2713  test_export_pdf - .pdf created",
], size=12, color=C_TEXT)
add_text(slide, 1.0, 6.2, 5, 0.4, "12/12 tests passing", size=14, color=C_INPUT, bold=True)

# Right: smoke test
add_shape(slide, 6.8, 1.3, 5.7, 5.5, C_PANEL)
add_text(slide, 7.0, 1.5, 5, 0.5, "Smoke Test", size=18, color=C_CHART, bold=True)
add_bullet_list(slide, 7.0, 2.2, 5.2, 4, [
    "\u2713  Load 2,000 rows x 7 columns",
    "\u2713  Filter 'North' \u2192 516 rows",
    "\u2713  Pivot table \u2192 4 rows",
    "\u2713  Describe \u2192 7 stat rows",
    "\u2713  Correlation \u2192 4x4 matrix",
    "\u2713  Regression R\u00b2 = 0.416",
    "\u2713  K-means \u2192 3 clusters",
    "\u2713  Chart drawn OK",
    "\u2713  Excel export: 6,249 bytes",
    "\u2713  PDF export: 21,931 bytes",
], size=12, color=C_TEXT)
add_text(slide, 7.0, 6.2, 5, 0.4, "ALL SMOKE TESTS PASSED", size=14, color=C_CHART, bold=True)

# ════════════════════════════════════════════════════════════════════════════
# SLIDE 9: Thank You
# ════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 1, 2.5, 11.3, 1.2, "Thank You", size=48, color=C_WHITE, bold=True,
         align=PPAlign.CENTER)
add_text(slide, 1, 3.8, 11.3, 0.8, "DataFlow Analyzer", size=24, color=C_CHART,
         align=PPAlign.CENTER)
features = [("CSV/Excel", C_INPUT), ("Filter & Pivot", C_PROC),
            ("Charts", C_CHART), ("Regression", C_ANAL), ("K-means", C_INPUT), ("Export", C_EXP)]
for i, (feat, color) in enumerate(features):
    fx = 1.5 + i * 1.8
    add_shape(slide, fx, 5.0, 1.6, 0.5, color)
    add_text(slide, fx, 5.05, 1.6, 0.4, feat, size=11, color=C_WHITE, bold=True,
             align=PPAlign.CENTER)
add_text(slide, 1, 6.2, 11.3, 0.5, "Built with Python, Tkinter, pandas & matplotlib",
         size=12, color=C_TEXT, align=PPAlign.CENTER)

# ── Save ──
path = os.path.join(OUT, "DataFlow_Analyzer_Presentation.pptx")
prs.save(path)
print(f"Presentation saved: {path}")
print(f"Slides: {len(prs.slides)}")
print(f"Size: {os.path.getsize(path):,} bytes")
